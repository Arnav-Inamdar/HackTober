from __future__ import annotations

import asyncio
from io import BytesIO

import pytest
import httpx
from fastapi.testclient import TestClient
from pydantic import ValidationError
from pptx import Presentation

from backend import main
from backend.ingestion import IngestionError, extract_document_text
from backend.knowledge import KnowledgeError, RetrievedKnowledge, KnowledgeSource, _CACHE, retrieve_wikipedia
from backend.ollama_client import (
    QUIZZES,
    StoredQuiz,
    normalize_model_questions,
    parse_json_response,
    role_system_prompt,
    store_quiz,
)
from backend.schemas import EvaluationRequest, GenerationRequest, QuestionSpec


@pytest.fixture
def client() -> TestClient:
    return TestClient(main.app)


def _request(**overrides):
    data = {
        "material": "Mitochondria convert energy into ATP through cellular respiration.",
        "topic": "Cell biology",
        "role": "strict_examiner",
        "custom_role": "",
        "difficulty": "medium",
        "count": 5,
        "style": "multiple_choice",
        "extra_instructions": "",
        "model": "installed-local-model",
    }
    data.update(overrides)
    return data


def _model_quiz_json() -> str:
    return '{"questions":[{"question":"What does cellular respiration produce?","type":"multiple_choice","options":["ATP","DNA","Oxygen","Protein"],"correct_answer":"ATP","explanation":"Mitochondria convert energy into ATP through cellular respiration.","topic":"Cell biology","difficulty":"medium"}]}'


def test_text_extraction_and_empty_file_errors():
    assert extract_document_text("notes.md", b"# Review\nATP stores usable energy.") == "# Review\nATP stores usable energy."
    with pytest.raises(IngestionError, match="empty"):
        extract_document_text("empty.txt", b"")
    with pytest.raises(IngestionError, match="not supported"):
        extract_document_text("notes.exe", b"no execution")


def test_pptx_extraction_preserves_slide_order():
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first_title = first.shapes.title
    assert first_title is not None
    first_title.text = "First topic"
    second = presentation.slides.add_slide(presentation.slide_layouts[0])
    second_title = second.shapes.title
    assert second_title is not None
    second_title.text = "Second topic"
    buffer = BytesIO()
    presentation.save(buffer)
    text = extract_document_text("ordered.pptx", buffer.getvalue())
    assert text.index("Slide 1") < text.index("First topic") < text.index("Slide 2") < text.index("Second topic")


def test_quiz_schema_accepts_only_requested_question_counts():
    for count in (5, 10, 15, 20, 30):
        assert GenerationRequest(**_request(count=count)).count == count
    with pytest.raises(ValidationError):
        GenerationRequest(**_request(count=6))


def test_question_normalization_deduplicates_and_checks_mc_options():
    request = GenerationRequest(**_request())
    payload = {"questions": [
        {"question": "Which molecule stores usable energy?", "type": "multiple_choice", "options": ["ATP", "DNA", "Oxygen"], "correct_answer": "ATP", "explanation": "ATP stores usable energy.", "topic": "Energy"},
        {"question": "Which molecule stores usable energy?", "type": "mcq", "options": ["ATP", "DNA", "Oxygen"], "correct_answer": "ATP", "explanation": "Duplicate.", "topic": "Energy"},
        {"question": "Which molecule is the answer?", "type": "multiple_choice", "options": ["DNA", "DNA", "Oxygen"], "correct_answer": "DNA", "explanation": "Duplicated choices are invalid.", "topic": "Energy"},
    ]}
    result = normalize_model_questions(payload, request)
    assert len(result) == 1
    assert result[0].correct_answer == "ATP"


def test_parse_json_fences_and_role_system_prompt():
    assert parse_json_response('```json\n{"ok": true}\n```') == {"ok": True}
    system = role_system_prompt("friendly_tutor")
    assert "encouraging, supportive" in system
    assert "safety rules above always take precedence" in system


def test_upload_endpoint_extracts_text_without_ollama(client: TestClient):
    response = client.post("/api/content/ingest", files=[("files", ("notes.txt", b"The Krebs cycle produces electron carriers.", "text/plain"))])
    assert response.status_code == 200
    body = response.json()
    assert body["files"][0]["status"] == "success"
    assert "Krebs cycle" in body["content"]


def test_health_reports_disconnected_ollama_as_recoverable(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    async def unavailable(_client, _base_url):
        raise main.OllamaError("offline")

    monkeypatch.setattr(main, "ollama_tags", unavailable)
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["ollama_connected"] is False
    assert response.json()["app"] == "ok"


def test_generate_hides_answer_and_evaluate_is_idempotent(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    async def tags(_client, _base_url):
        return [{"name": "installed-local-model"}]

    async def generate(_client, _base_url, _model, _prompt, *, system_prompt=None):
        assert system_prompt and "exacting and challenging" in system_prompt
        return _model_quiz_json()

    monkeypatch.setattr(main, "ollama_tags", tags)
    monkeypatch.setattr(main, "call_generate", generate)
    QUIZZES.clear()
    response = client.post("/api/quiz/generate", json=_request())
    assert response.status_code == 200
    generated = response.json()
    assert generated["warning"] == "The model provided 1 valid questions of the 5 requested."
    assert "correct_answer" not in generated["questions"][0]
    assert "explanation" not in generated["questions"][0]

    evaluate_url = "/api/quiz/evaluate"
    wrong = client.post(evaluate_url, json={"quiz_id": generated["quiz_id"], "question_id": "q-1", "answer": "Oxygen"})
    replay = client.post(evaluate_url, json={"quiz_id": generated["quiz_id"], "question_id": "q-1", "answer": "ATP"})
    assert wrong.status_code == replay.status_code == 200
    assert wrong.json()["status"] == "incorrect"
    assert replay.json() == wrong.json()
    assert wrong.json()["expected_answer"] == "ATP"
    assert "Mitochondria convert energy" in wrong.json()["explanation"]


def test_validation_error_does_not_echo_submitted_material(client: TestClient):
    response = client.post("/api/quiz/generate", json=_request(material="PRIVATE-STUDY-NOTE"))
    assert response.status_code == 422
    assert "PRIVATE-STUDY-NOTE" not in response.text
    assert "at least 20 characters" in response.text


def test_concurrent_short_answer_submission_uses_one_cached_evaluation(monkeypatch: pytest.MonkeyPatch):
    calls = 0

    class DummyClient:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return False

    async def evaluate_once(_client, _base_url, _model, _prompt, *, system_prompt=None):
        nonlocal calls
        calls += 1
        await asyncio.sleep(0.02)
        return '{"status":"correct","feedback":"The essential concept is present."}'

    monkeypatch.setattr(main, "create_http_client", DummyClient)
    monkeypatch.setattr(main, "call_generate", evaluate_once)
    QUIZZES.clear()
    quiz_id = store_quiz(StoredQuiz(
        model="local-model",
        role="friendly_tutor",
        custom_role="",
        material="Photosynthesis stores light energy as chemical energy.",
        questions=[QuestionSpec(
            question="What energy conversion happens in photosynthesis?",
            type="short_answer",
            options=[],
            correct_answer="Light energy is stored as chemical energy.",
            explanation="Photosynthesis converts light energy into chemical energy.",
            topic="Photosynthesis",
            difficulty="medium",
        )],
    ))

    async def submit_together():
        return await asyncio.gather(
            main.evaluate_answer(EvaluationRequest(quiz_id=quiz_id, question_id="q-1", answer="Equivalent wording A")),
            main.evaluate_answer(EvaluationRequest(quiz_id=quiz_id, question_id="q-1", answer="Different wording B")),
        )

    first, second = asyncio.run(submit_together())
    assert calls == 1
    assert first.model_dump() == second.model_dump()
    assert first.expected_answer == "Light energy is stored as chemical energy."


def test_source_modes_require_only_the_selected_input():
    topic_only = GenerationRequest(**_request(source_mode="live_topic", material="", topic="AI"))
    assert topic_only.material == ""
    assert topic_only.source_mode == "live_topic"
    with pytest.raises(ValidationError, match="at least 20 characters"):
        GenerationRequest(**_request(source_mode="my_material", material="", topic="AI"))


def test_wikipedia_retrieval_uses_batched_plaintext_and_local_cache():
    calls = []
    article_text = "Photosynthesis converts light energy into chemical energy. " * 9

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        params = request.url.params
        assert request.headers["user-agent"].startswith("AVIA-LocalQuiz/")
        if params.get("list") == "search":
            assert params.get("srsearch") == "Photosynthesis"
            assert params.get("srnamespace") == "0"
            assert params.get("srlimit") == "5"
            return httpx.Response(200, json={"query": {"search": [{"ns": 0, "pageid": 17, "title": "Photosynthesis"}]}})
        assert params.get("prop") == "extracts|info"
        assert params.get("explaintext") == "1"
        assert params.get("pageids") == "17"
        return httpx.Response(200, json={"query": {"pages": [{"ns": 0, "pageid": 17, "title": "Photosynthesis", "fullurl": "https://en.wikipedia.org/wiki/Photosynthesis", "extract": article_text}]}})

    async def run():
        _CACHE.clear()
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            first = await retrieve_wikipedia(client, "Photosynthesis")
            second = await retrieve_wikipedia(client, "  photosynthesis  ")
        return first, second

    first, second = asyncio.run(run())
    assert len(calls) == 2  # One search and one batched extract request; the second call is cached.
    assert first == second
    assert "chemical energy" in first.context
    assert first.sources == (KnowledgeSource(title="Photosynthesis", url="https://en.wikipedia.org/wiki/Photosynthesis"),)


def test_wikipedia_rate_limit_has_a_safe_recovery_message():
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda _request: httpx.Response(429))) as client:
            return await retrieve_wikipedia(client, "A unique rate-limit test topic")

    with pytest.raises(KnowledgeError, match="rate limiting"):
        asyncio.run(run())


def test_topic_generation_returns_citations_and_never_uses_private_notes_as_retrieval_input(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    async def tags(_client, _base_url):
        return [{"name": "installed-local-model"}]

    async def retrieval(_client, topic):
        assert topic == "Photosynthesis"
        return RetrievedKnowledge(
            topic=topic,
            context="[Retrieved source: Photosynthesis]\nLight energy becomes chemical energy in photosynthesis.",
            sources=(KnowledgeSource(title="Photosynthesis", url="https://en.wikipedia.org/wiki/Photosynthesis"),),
        )

    async def generate(_client, _base_url, _model, prompt, *, system_prompt=None):
        assert "Light energy becomes chemical energy" in prompt
        assert "PRIVATE-UPLOADED-NOTE" not in prompt
        assert system_prompt is not None
        return _model_quiz_json()

    monkeypatch.setattr(main, "ollama_tags", tags)
    monkeypatch.setattr(main, "retrieve_wikipedia", retrieval)
    monkeypatch.setattr(main, "call_generate", generate)
    QUIZZES.clear()
    response = client.post("/api/quiz/generate", json=_request(
        source_mode="live_topic",
        material="PRIVATE-UPLOADED-NOTE must not be sent to public search or used instead of retrieved evidence.",
        topic="Photosynthesis",
    ))
    assert response.status_code == 200
    payload = response.json()
    assert payload["sources"] == [{"title": "Photosynthesis", "url": "https://en.wikipedia.org/wiki/Photosynthesis"}]
    assert QUIZZES[payload["quiz_id"]].material.startswith("[Retrieved source: Photosynthesis]")


def test_material_generation_does_not_retrieve_public_sources(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    async def tags(_client, _base_url):
        return [{"name": "installed-local-model"}]

    async def retrieval(*_args):
        raise AssertionError("material-only mode must not call public search")

    async def generate(_client, _base_url, _model, _prompt, *, system_prompt=None):
        return _model_quiz_json()

    monkeypatch.setattr(main, "ollama_tags", tags)
    monkeypatch.setattr(main, "retrieve_wikipedia", retrieval)
    monkeypatch.setattr(main, "call_generate", generate)
    QUIZZES.clear()
    response = client.post("/api/quiz/generate", json=_request(source_mode="my_material", topic="Cell biology"))
    assert response.status_code == 200
    assert response.json()["sources"] == []


def test_wikipedia_no_results_offers_a_title_or_notes_fallback():
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda _request: httpx.Response(200, json={"query": {"search": []}}))) as client:
            return await retrieve_wikipedia(client, "No results AVIA test phrase")

    with pytest.raises(KnowledgeError, match="no clear article"):
        asyncio.run(run())


def test_wikipedia_timeout_offers_a_retry_or_material_fallback():
    def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ReadTimeout("upstream timeout", request=request)

    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await retrieve_wikipedia(client, "Timeout AVIA test phrase")

    with pytest.raises(KnowledgeError, match="timed out"):
        asyncio.run(run())
