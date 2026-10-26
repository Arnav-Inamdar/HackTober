from __future__ import annotations

from io import BytesIO
from pathlib import Path

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_FILES = 10
MAX_TOTAL_BYTES = 40 * 1024 * 1024
MAX_CONTEXT_CHARS = 80_000
SUPPORTED_TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}
SUPPORTED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED_EXTENSIONS = SUPPORTED_TEXT_EXTENSIONS | SUPPORTED_IMAGE_EXTENSIONS | {".pdf", ".pptx"}


class IngestionError(ValueError):
    """A user-safe error describing why one upload could not be processed."""


def extract_document_text(filename: str, data: bytes) -> str:
    """Extract text from supported non-image formats without writing uploads to disk."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise IngestionError("This file type is not supported. Use TXT, Markdown, PDF, PPTX, or an image.")
    if not data:
        raise IngestionError("This file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise IngestionError("This file is larger than the 15 MB limit.")

    if ext in SUPPORTED_TEXT_EXTENSIONS:
        try:
            text = data.decode("utf-8-sig").strip()
        except UnicodeDecodeError as exc:
            raise IngestionError("The text file is not readable UTF-8. Save it as UTF-8 and try again.") from exc
    elif ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data), strict=False)
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n\n".join(f"Page {idx + 1}\n{page.strip()}" for idx, page in enumerate(pages) if page.strip())
        except Exception as exc:
            raise IngestionError("This PDF is damaged, encrypted, or could not be read locally.") from exc
    elif ext == ".pptx":
        try:
            from pptx import Presentation
            deck = Presentation(BytesIO(data))
            slide_text: list[str] = []
            for idx, slide in enumerate(deck.slides, start=1):
                parts: list[str] = []
                title_shape = slide.shapes.title
                if title_shape is not None and title_shape.text.strip():
                    parts.append(title_shape.text.strip())
                for shape in slide.shapes:
                    if shape is title_shape:
                        continue
                    shape_text = getattr(shape, "text", "")
                    if isinstance(shape_text, str) and shape_text.strip():
                        parts.append(shape_text.strip())
                    table = getattr(shape, "table", None)
                    if table is not None:
                        for row in getattr(table, "rows", ()):
                            cells = [cell_text for cell in row.cells if (cell_text := str(getattr(cell, "text", "")).strip())]
                            if cells:
                                parts.append(" | ".join(cells))
                if parts:
                    slide_text.append(f"Slide {idx}\n" + "\n".join(parts))
            text = "\n\n".join(slide_text)
        except ImportError as exc:
            raise IngestionError("PPTX support is unavailable. Install the project dependencies and retry.") from exc
        except Exception as exc:
            raise IngestionError("This presentation is damaged or could not be read locally.") from exc
    else:
        raise IngestionError("Image files must be interpreted by a compatible local vision model or OCR.")

    if not text.strip():
        raise IngestionError("No readable text was found. Try a text-based file or use local OCR for scanned pages.")
    return text.strip()[:MAX_CONTEXT_CHARS]


def validate_image(data: bytes) -> None:
    try:
        from PIL import Image
        with Image.open(BytesIO(data)) as image:
            image.verify()
    except ImportError as exc:
        raise IngestionError("Image validation support is unavailable in this Python environment.") from exc
    except Exception as exc:
        raise IngestionError("The image is corrupted or unreadable.") from exc


def local_ocr(data: bytes) -> str:
    try:
        from PIL import Image
        import pytesseract
        text = pytesseract.image_to_string(Image.open(BytesIO(data))).strip()
    except ImportError as exc:
        raise IngestionError("No compatible local vision model is selected, and the optional OCR package is unavailable.") from exc
    except Exception as exc:
        raise IngestionError("Local OCR could not read this image. Check that Tesseract OCR is installed, or select a vision-capable Ollama model.") from exc
    if not text:
        raise IngestionError("Local OCR found no readable text in this image.")
    return text[:MAX_CONTEXT_CHARS]
